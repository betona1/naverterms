import io
import os
import csv
import sys
import subprocess
import tempfile
import zipfile
import logging
from datetime import datetime

import openpyxl
from django.conf import settings
from django.db import connections

logger = logging.getLogger(__name__)

DB = 'ads'

# 엑셀 컬럼 인덱스 (0-based) → DB 필드명 매핑
EXCEL_COL_MAP = {
    0: 'seller_code1',
    1: 'seller_code2',
    # 2 = product_code (별도 처리)
    3: 'category_code',
    4: 'category_name',
    5: 'market_category',
    6: 'product_name',
    7: 'market_product_name',
    8: 'ownerclan_price',
    9: 'consumer_price',
    10: 'market_price',
    11: 'shipping_fee',
    12: 'shipping_type',
    13: 'min_qty',
    14: 'max_qty',
    15: 'company_notice',
    16: 'special_notice',
    17: 'option1_name',
    18: 'option1_values',
    19: 'option2_name',
    20: 'option2_values',
    21: 'combined_option',
    22: 'product_attribute',
    23: 'product_grade',
    24: 'tax_type',
    25: 'compliance',
    26: 'age_restriction',
    27: 'return_possible',
    28: 'image_large',
    29: 'image_medium',
    30: 'image_small',
    31: 'manufacturer',
    32: 'brand',
    33: 'model_name',
    34: 'origin',
    35: 'keywords',
    36: 'registered_at',
    37: 'modified_at',
    38: 'header_text',
    39: 'detail_html',
    40: 'notice_code',
    41: 'notice_category',
    42: 'notice_info',
    43: 'notice_html',
    44: 'market_gmarket',
    45: 'market_auction',
    46: 'market_11st',
    47: 'market_coupang',
    48: 'market_smartstore',
    49: 'market_promo',
    50: 'market_gift',
    51: 'certification_type',
    52: 'certification_info',
    53: 'return_fee',
    54: 'independent_option',
    55: 'combined_option_detail',
}

INT_FIELDS = {
    'ownerclan_price', 'consumer_price', 'market_price',
    'shipping_fee', 'min_qty', 'max_qty', 'return_fee',
}

DATETIME_FIELDS = {'registered_at', 'modified_at'}

TRACKABLE_FIELDS = list(EXCEL_COL_MAP.values())

# 수정사항 감지용 핵심 필드 (orig_ 비교)
CHANGE_DETECT_FIELDS = {
    'ownerclan_price': 'int', 'consumer_price': 'int', 'market_price': 'int',
    'product_name': 'str', 'market_product_name': 'str',
    'detail_html': 'str', 'header_text': 'str',
    'image_large': 'str', 'image_medium': 'str', 'image_small': 'str',
    'shipping_fee': 'int', 'shipping_type': 'str', 'return_fee': 'int',
    'combined_option': 'str', 'independent_option': 'str', 'combined_option_detail': 'str',
    'option1_name': 'str', 'option1_values': 'str', 'option2_name': 'str', 'option2_values': 'str',
    'manufacturer': 'str', 'origin': 'str', 'brand': 'str', 'model_name': 'str',
    'product_attribute': 'str',
    'keywords': 'str',
    'compliance': 'str', 'age_restriction': 'str', 'return_possible': 'str',
    'certification_type': 'str', 'certification_info': 'str',
    'company_notice': 'str', 'special_notice': 'str',
    'notice_code': 'str', 'notice_category': 'str', 'notice_info': 'str', 'notice_html': 'str',
}

# 필드 → 변경 그룹 매핑
FIELD_TO_GROUP = {
    'ownerclan_price': 'price', 'consumer_price': 'price', 'market_price': 'price',
    'shipping_fee': 'shipping', 'shipping_type': 'shipping', 'return_fee': 'shipping',
    'product_name': 'product_name', 'market_product_name': 'product_name',
    'detail_html': 'detail', 'header_text': 'detail', 'keywords': 'detail',
    'image_large': 'image', 'image_medium': 'image', 'image_small': 'image',
    'combined_option': 'option', 'independent_option': 'option', 'combined_option_detail': 'option',
    'option1_name': 'option', 'option1_values': 'option', 'option2_name': 'option', 'option2_values': 'option',
    'manufacturer': 'info', 'brand': 'info', 'model_name': 'info', 'origin': 'info', 'product_attribute': 'info',
    'compliance': 'compliance', 'age_restriction': 'compliance', 'return_possible': 'compliance',
    'certification_type': 'compliance', 'certification_info': 'compliance',
    'company_notice': 'notice', 'special_notice': 'notice',
    'notice_code': 'notice', 'notice_category': 'notice', 'notice_info': 'notice', 'notice_html': 'notice',
}

CHANGE_GROUP_LABELS = {
    'price': '가격변동', 'shipping': '배송비변동', 'product_name': '상품명변경',
    'detail': '상세페이지', 'image': '이미지변경', 'option': '옵션변경',
    'info': '제품정보', 'compliance': '인증/반품', 'notice': '공지변경',
}


def _any_field_changed_sql():
    """수정사항이 있는 상품 필터 SQL 조건"""
    parts = []
    for f, t in CHANGE_DETECT_FIELDS.items():
        if t == 'int':
            parts.append(f"COALESCE({f},0) != COALESCE(orig_{f},0)")
        else:
            parts.append(f"COALESCE({f},'') != COALESCE(orig_{f},'')")
    return '(' + ' OR '.join(parts) + ')'


def _dictfetchall(cursor):
    columns = [col[0] for col in cursor.description]
    return [dict(zip(columns, row)) for row in cursor.fetchall()]


def _safe_str(val):
    if val is None:
        return ''
    return str(val).strip()


def _safe_int(val):
    if val is None:
        return 0
    try:
        return int(float(val))
    except (ValueError, TypeError):
        return 0


def _safe_datetime(val):
    if val is None or val == '':
        return None
    if isinstance(val, datetime):
        return val
    s = str(val).strip()
    for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%Y-%m-%d'):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def _parse_excel_row(row_values):
    data = {}
    for col_idx, field_name in EXCEL_COL_MAP.items():
        raw = row_values[col_idx] if col_idx < len(row_values) else None
        if field_name in INT_FIELDS:
            data[field_name] = _safe_int(raw)
        elif field_name in DATETIME_FIELDS:
            data[field_name] = _safe_datetime(raw)
        else:
            data[field_name] = _safe_str(raw)
    return data


def _field_changed(old_val, new_val, field_name):
    if field_name in INT_FIELDS:
        return _safe_int(old_val) != _safe_int(new_val)
    if field_name in DATETIME_FIELDS:
        a = _safe_datetime(old_val)
        b = _safe_datetime(new_val)
        if a and b:
            return a.strftime('%Y-%m-%d %H:%M:%S') != b.strftime('%Y-%m-%d %H:%M:%S')
        return (a is None) != (b is None)
    return _safe_str(old_val) != _safe_str(new_val)


def upload_excel_async(uploaded_file):
    suffix = '.zip' if uploaded_file.name.lower().endswith('.zip') else '.xlsx'
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix, prefix='ownerclan_')
    for chunk in uploaded_file.chunks():
        tmp.write(chunk)
    tmp.close()

    now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    with connections[DB].cursor() as cur:
        cur.execute(
            "INSERT INTO lohas_task (task_type, status, input_data, result_data, created_at, updated_at) "
            "VALUES (%s, %s, %s, %s, %s, %s)",
            ['ownerclan_upload', 'pending',
             '{"file_path": "%s", "filename": "%s"}' % (tmp.name, uploaded_file.name),
             '{}', now, now],
        )
        task_id = cur.lastrowid

    manage_py = os.path.join(settings.BASE_DIR, 'manage.py')
    proc = subprocess.Popen(
        [sys.executable, manage_py, 'ownerclan_upload', str(task_id)],
        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
    )

    with connections[DB].cursor() as cur:
        cur.execute("UPDATE lohas_task SET pid=%s WHERE id=%s", [proc.pid, task_id])

    return {'task_id': task_id, 'status': 'pending'}


def get_upload_task(task_id):
    with connections[DB].cursor() as cur:
        cur.execute(
            "SELECT id, task_type, status, input_data, result_data, pid, created_at, updated_at "
            "FROM lohas_task WHERE id=%s", [task_id])
        rows = _dictfetchall(cur)
        if not rows:
            return None
        row = rows[0]
        # result_data / input_data 는 JSON 텍스트일 수 있음
        import json
        for k in ('input_data', 'result_data'):
            if isinstance(row[k], str):
                try:
                    row[k] = json.loads(row[k])
                except (json.JSONDecodeError, TypeError):
                    pass
        return row


def check_running_task():
    with connections[DB].cursor() as cur:
        cur.execute(
            "SELECT id FROM lohas_task WHERE task_type='ownerclan_upload' AND status IN ('pending','running') LIMIT 1"
        )
        row = cur.fetchone()
        return row[0] if row else None


def upload_csv_status(uploaded_file):
    content = uploaded_file.read()
    for enc in ('euc-kr', 'utf-8', 'cp949'):
        try:
            text = content.decode(enc, errors='replace')
            break
        except Exception:
            continue
    else:
        text = content.decode('utf-8', errors='replace')

    reader = csv.reader(io.StringIO(text))
    next(reader, None)
    next(reader, None)

    STATUS_MAP = {
        '품절': 2, '단종': 3, '유통금지': 3,
        '옵션 품절': 2, '옵션 단종': 3,
        '재입고': 1, '옵션 재입고': 1,
    }

    updated = 0
    with connections[DB].cursor() as cur:
        for row in reader:
            if len(row) < 8:
                continue
            code = row[1].strip()
            status_text = row[7].strip()
            if not code:
                continue
            sale_status = STATUS_MAP.get(status_text)
            if sale_status is not None:
                cur.execute(
                    "UPDATE ownerclan_product SET sale_status=%s WHERE product_code=%s",
                    [sale_status, code],
                )
                if cur.rowcount > 0:
                    updated += 1

    return {'updated': updated}


def sync_products(product_ids=None):
    fields = list(EXCEL_COL_MAP.values())
    set_parts = [f"orig_{f} = {f}" for f in fields]
    set_parts.append("is_synced = 1")
    set_parts.append("synced_at = %s")
    now = datetime.now()

    with connections[DB].cursor() as cur:
        if product_ids:
            placeholders = ','.join(['%s'] * len(product_ids))
            cur.execute(
                f"UPDATE ownerclan_product SET {', '.join(set_parts)} "
                f"WHERE id IN ({placeholders})",
                [now] + product_ids,
            )
        else:
            cur.execute(
                f"UPDATE ownerclan_product SET {', '.join(set_parts)} "
                f"WHERE is_synced = 0",
                [now],
            )
        count = cur.rowcount
    return {'synced': count}


def _changed_field_condition(field_name):
    if field_name in INT_FIELDS:
        return f"COALESCE({field_name},0) != COALESCE(orig_{field_name},0)"
    elif field_name in DATETIME_FIELDS:
        return f"COALESCE(CAST({field_name} AS CHAR),'') != COALESCE(CAST(orig_{field_name} AS CHAR),'')"
    else:
        return f"COALESCE({field_name},'') != COALESCE(orig_{field_name},'')"


def _build_where(sale_status=None, is_synced=None, search=None, changed_field=None):
    where = ['1=1']
    params = []
    if sale_status is not None:
        sale_str = str(sale_status)
        if ',' in sale_str:
            vals = [int(v.strip()) for v in sale_str.split(',') if v.strip()]
            ph = ','.join(['%s'] * len(vals))
            where.append(f'sale_status IN ({ph})')
            params.extend(vals)
        else:
            where.append('sale_status = %s')
            params.append(int(sale_status))
    if is_synced is not None:
        where.append('is_synced = %s')
        params.append(int(is_synced))
    if search:
        import re
        codes = [c.strip() for c in re.split(r'[\s,\n\r\t]+', search) if c.strip()]
        if len(codes) > 1:
            ph = ','.join(['%s'] * len(codes))
            where.append(f'product_code IN ({ph})')
            params.extend(codes)
        else:
            where.append('(product_code LIKE %s OR product_name LIKE %s)')
            like = f'%{search}%'
            params.extend([like, like])
    if changed_field:
        if changed_field in TRACKABLE_FIELDS:
            where.append(_changed_field_condition(changed_field))
        elif changed_field == '__any__':
            where.append(_any_field_changed_sql())
    return ' AND '.join(where), params


def get_products(page=1, per_page=50, sale_status=None, is_synced=None, search=None, changed_field=None):
    where_sql, params = _build_where(sale_status, is_synced, search, changed_field)
    offset = (page - 1) * per_page

    with connections[DB].cursor() as cur:
        cur.execute(f"SELECT COUNT(*) FROM ownerclan_product WHERE {where_sql}", params)
        total = cur.fetchone()[0]

        changed_sql = _any_field_changed_sql()
        cur.execute(
            f"SELECT id, product_code, product_name, orig_product_name, "
            f"market_product_name, orig_market_product_name, "
            f"ownerclan_price, orig_ownerclan_price, "
            f"market_price, orig_market_price, "
            f"shipping_fee, orig_shipping_fee, "
            f"return_fee, orig_return_fee, "
            f"image_large, orig_image_large, "
            f"image_small, "
            f"sale_status, is_synced, "
            f"category_name, manufacturer, origin, "
            f"uploaded_at, synced_at, created_at, "
            f"{changed_sql} as has_changes "
            f"FROM ownerclan_product WHERE {where_sql} "
            f"ORDER BY product_code "
            f"LIMIT %s OFFSET %s",
            params + [per_page, offset],
        )
        rows = _dictfetchall(cur)

    for r in rows:
        for k in ('uploaded_at', 'synced_at', 'created_at'):
            if r.get(k) and isinstance(r[k], datetime):
                r[k] = r[k].isoformat()

    return {
        'items': rows,
        'total': total,
        'page': page,
        'per_page': per_page,
        'total_pages': (total + per_page - 1) // per_page if total > 0 else 0,
    }


def get_product_detail(product_id):
    fields = list(EXCEL_COL_MAP.values())
    select_parts = ['id', 'product_code', 'sale_status', 'is_synced',
                     'synced_at', 'uploaded_at', 'created_at', 'updated_at']
    for f in fields:
        select_parts.append(f)
        select_parts.append(f'orig_{f}')

    with connections[DB].cursor() as cur:
        cur.execute(
            f"SELECT {', '.join(select_parts)} FROM ownerclan_product WHERE id=%s",
            [product_id],
        )
        rows = _dictfetchall(cur)
        if not rows:
            return None
        row = rows[0]

    for k, v in row.items():
        if isinstance(v, datetime):
            row[k] = v.isoformat()

    changed_fields = []
    for f in fields:
        if _field_changed(row.get(f'orig_{f}'), row.get(f), f):
            changed_fields.append(f)
    row['changed_fields'] = changed_fields
    return row


def get_stats():
    changed_sql = _any_field_changed_sql()
    with connections[DB].cursor() as cur:
        cur.execute(
            "SELECT "
            "COUNT(*) as total, "
            "SUM(sale_status=1) as selling, "
            "SUM(sale_status=2) as soldout, "
            "SUM(sale_status=3) as discontinued, "
            f"SUM({changed_sql}) as changed "
            "FROM ownerclan_product"
        )
        row = _dictfetchall(cur)[0]
    return {k: int(v or 0) for k, v in row.items()}


def get_changed_field_counts():
    counts = {}
    with connections[DB].cursor() as cur:
        for f in TRACKABLE_FIELDS:
            cond = _changed_field_condition(f)
            cur.execute(f"SELECT COUNT(*) FROM ownerclan_product WHERE {cond}")
            cnt = cur.fetchone()[0]
            if cnt > 0:
                counts[f] = cnt
    return counts


def get_products_for_export(sale_status=None, is_synced=None, search=None, changed_field=None):
    where_sql, params = _build_where(sale_status, is_synced, search, changed_field)

    fields = list(EXCEL_COL_MAP.values())
    orig_fields = [f'orig_{f}' for f in fields]

    select_cols = (
        ['product_code', 'sale_status', 'is_synced']
        + fields + orig_fields
        + ['uploaded_at', 'synced_at']
    )

    with connections[DB].cursor() as cur:
        cur.execute(
            f"SELECT {', '.join(select_cols)} "
            f"FROM ownerclan_product WHERE {where_sql} "
            f"ORDER BY product_code",
            params,
        )
        return _dictfetchall(cur)


def upload_soldout_txt(uploaded_file):
    content = uploaded_file.read()
    for enc in ('utf-8', 'euc-kr', 'cp949'):
        try:
            text = content.decode(enc)
            break
        except Exception:
            continue
    else:
        text = content.decode('utf-8', errors='replace')

    soldout_codes = {line.strip() for line in text.splitlines() if line.strip()}

    with connections[DB].cursor() as cur:
        cur.execute("SELECT product_code FROM ownerclan_product")
        db_codes = {r[0] for r in cur.fetchall()}

    matched = soldout_codes & db_codes
    today = datetime.now().strftime('%Y-%m-%d')

    ownerclan_updated = 0
    batch = list(matched)
    with connections[DB].cursor() as cur:
        for i in range(0, len(batch), 1000):
            chunk = batch[i:i + 1000]
            placeholders = ','.join(['%s'] * len(chunk))
            cur.execute(
                f"UPDATE ownerclan_product SET sale_status=2, modified_at=%s "
                f"WHERE product_code IN ({placeholders}) AND sale_status != 2",
                [today] + chunk,
            )
            ownerclan_updated += cur.rowcount

    smartstore_updated = 0
    with connections['myproduct'].cursor() as cur:
        cur.execute("UPDATE smartstore_product SET ownerclan_soldout=0 WHERE ownerclan_soldout=1")
        for i in range(0, len(batch), 1000):
            chunk = batch[i:i + 1000]
            placeholders = ','.join(['%s'] * len(chunk))
            cur.execute(
                f"UPDATE smartstore_product SET ownerclan_soldout=1 "
                f"WHERE seller_management_code IN ({placeholders})",
                chunk,
            )
            smartstore_updated += cur.rowcount

    return {
        'total_codes': len(soldout_codes),
        'ownerclan_matched': len(matched),
        'ownerclan_updated': ownerclan_updated,
        'smartstore_updated': smartstore_updated,
    }


def activate_suspended_from_task(task_id):
    """업로드 태스크에 포함된 상품 중 판매중지인 것을 판매중(1)으로 변경"""
    import json as _json

    with connections[DB].cursor() as cur:
        cur.execute("SELECT input_data FROM lohas_task WHERE id=%s", [task_id])
        row = cur.fetchone()
        if not row:
            return {'error': 'task not found', 'activated': 0}

        input_data = row[0]
        if isinstance(input_data, str):
            input_data = _json.loads(input_data)

        file_path = input_data.get('file_path', '')
        filename = input_data.get('filename', '')

    # 태스크의 업로드 파일에서 product_code 목록을 다시 읽을 수 없으므로
    # (파일은 이미 삭제됨) DB에서 최근 업로드된 상품 중 판매중지인 것을 직접 갱신
    # uploaded_at이 이 태스크의 시작 이후인 상품 = 이 태스크에서 처리된 상품
    with connections[DB].cursor() as cur:
        cur.execute("SELECT created_at FROM lohas_task WHERE id=%s", [task_id])
        task_row = cur.fetchone()
        if not task_row:
            return {'activated': 0}
        task_created = task_row[0]

        cur.execute(
            "UPDATE ownerclan_product SET sale_status=1 "
            "WHERE sale_status != 1 AND uploaded_at >= %s",
            [task_created],
        )
        activated = cur.rowcount

    return {'activated': activated}


def get_w_codes(sale_status=None, is_synced=None, search=None, changed_field=None):
    where_sql, params = _build_where(sale_status, is_synced, search, changed_field)
    with connections[DB].cursor() as cur:
        cur.execute(
            f"SELECT product_code FROM ownerclan_product "
            f"WHERE {where_sql} ORDER BY product_code",
            params,
        )
        return [row[0] for row in cur.fetchall()]


def sync_status_preview():
    """스마트스토어 기준 동기화 미리보기 — 오너클랜 vs 스마트스토어 상태 비교"""
    with connections[DB].cursor() as cur:
        # 스마트스토어 SALE → 품절처리 (오너클랜 품절/단종)
        cur.execute('''
            SELECT op.sale_status, COUNT(*)
            FROM ownerclan_product op
            JOIN myproduct.smartstore_product sp ON sp.seller_management_code = op.product_code
            WHERE op.sale_status IN (2, 3) AND sp.status_type = 'SALE'
            GROUP BY op.sale_status
        ''')
        to_suspend = {}
        for row in cur.fetchall():
            to_suspend[int(row[0])] = row[1]

        # 스마트스토어 품절/중지 → SALE (오너클랜 판매중)
        cur.execute('''
            SELECT sp.status_type, COUNT(*)
            FROM ownerclan_product op
            JOIN myproduct.smartstore_product sp ON sp.seller_management_code = op.product_code
            WHERE op.sale_status = 1 AND sp.status_type IN ('SUSPENSION', 'OUTOFSTOCK')
            GROUP BY sp.status_type
        ''')
        to_activate = {}
        for row in cur.fetchall():
            to_activate[row[0]] = row[1]
        to_activate_total = sum(to_activate.values())

        # 필드별 변경 건수
        field_changes = {}
        field_queries = {
            'price': "ownerclan_price != orig_ownerclan_price OR market_price != orig_market_price",
            'product_name': "product_name != orig_product_name OR market_product_name != orig_market_product_name",
            'detail_html': "detail_html != orig_detail_html",
            'image': "image_large != orig_image_large",
            'shipping_fee': "shipping_fee != orig_shipping_fee",
            'return_fee': "return_fee != orig_return_fee",
            'combined_option': "combined_option != orig_combined_option",
            'origin': "origin != orig_origin",
            'manufacturer': "manufacturer != orig_manufacturer",
        }
        for key, cond in field_queries.items():
            cur.execute(f'SELECT COUNT(*) FROM ownerclan_product WHERE {cond}')
            cnt = cur.fetchone()[0]
            if cnt > 0:
                field_changes[key] = cnt

    return {
        'to_soldout': to_suspend.get(2, 0),
        'to_discontinued': to_suspend.get(3, 0),
        'to_activate_total': to_activate_total,
        'to_activate_suspension': to_activate.get('SUSPENSION', 0),
        'to_activate_outofstock': to_activate.get('OUTOFSTOCK', 0),
        'price_changed': field_changes.get('price', 0),
        'field_changes': field_changes,
        'total_sync': to_suspend.get(2, 0) + to_suspend.get(3, 0) + to_activate_total,
    }


_SYNC_LOCK_PATH = os.path.join(tempfile.gettempdir(), 'sync_ownerclan_status.lock')


def _sync_store_batch(store_id, items, api_key, secret_key):
    """단일 상점의 상품들을 상태 변경 (스레드 워커).
    returns: {'success': int, 'skipped': int, 'errors': [], 'skipped_items': [], 'changes': []}
    """
    import time
    import requests
    from smartstore.smartstore_product_service import _get_access_token

    res = {'success': 0, 'skipped': 0, 'errors': [], 'skipped_items': [], 'changes': []}

    if not api_key or not secret_key:
        for t, ns in items:
            res['errors'].append({'product_code': t['product_code'], 'error': 'API 키 미등록'})
        return res

    try:
        token = _get_access_token(api_key, secret_key)
    except Exception as e:
        for t, ns in items:
            res['errors'].append({'product_code': t['product_code'], 'error': f'토큰 실패: {e}'})
        return res

    headers = {'Authorization': f'Bearer {token}', 'Content-Type': 'application/json'}

    for target, new_status in items:
        opno = target['origin_product_no']
        old_ss = target.get('old_ss_status', '')
        url = f'https://api.commerce.naver.com/external/v2/products/origin-products/{opno}'

        for attempt in range(4):
            try:
                get_resp = requests.get(url, headers=headers, timeout=30)
                get_resp.raise_for_status()
                product = get_resp.json()['originProduct']

                cur_status = product.get('statusType', '')
                if cur_status and cur_status != new_status:
                    if cur_status == new_status:
                        pass
                    elif new_status == 'SALE' and cur_status in ('PROHIBITION', 'CLOSE', 'UNADMISSION'):
                        with connections['myproduct'].cursor() as c:
                            c.execute("UPDATE smartstore_product SET status_type=%s WHERE id=%s", [cur_status, target['id']])
                        if old_ss != cur_status:
                            res['changes'].append((target, old_ss, cur_status))
                        res['skipped'] += 1
                        res['skipped_items'].append({'product_code': target['product_code'], 'status': cur_status, 'store': target.get('store_name', '')})
                        time.sleep(0.3)
                        break
                    elif new_status == 'SUSPENSION' and cur_status in ('SUSPENSION', 'OUTOFSTOCK'):
                        with connections['myproduct'].cursor() as c:
                            c.execute("UPDATE smartstore_product SET status_type=%s WHERE id=%s", [cur_status, target['id']])
                        if old_ss != cur_status:
                            res['changes'].append((target, old_ss, cur_status))
                        res['success'] += 1
                        time.sleep(0.3)
                        break

                product['statusType'] = new_status
                put_resp = requests.put(url, json={'originProduct': product}, headers=headers, timeout=30)
                put_resp.raise_for_status()

                with connections['myproduct'].cursor() as c:
                    c.execute("UPDATE smartstore_product SET status_type=%s WHERE id=%s", [new_status, target['id']])
                res['changes'].append((target, old_ss, new_status))
                res['success'] += 1
                logger.info(f'[SYNC] {target["product_code"]}: {old_ss}->{new_status} ({target.get("store_name","")})')
                time.sleep(0.5)
                break

            except requests.exceptions.HTTPError as e:
                sc = e.response.status_code if e.response is not None else 0
                if sc in (401, 429):
                    wait = max(int(e.response.headers.get('Retry-After', 10)), 5 * (attempt + 1)) if sc == 429 else 2
                    time.sleep(wait)
                    try:
                        token = _get_access_token(api_key, secret_key)
                        headers = {'Authorization': f'Bearer {token}', 'Content-Type': 'application/json'}
                    except Exception:
                        pass
                    continue
                elif sc in (400, 404):
                    db_st = 'CLOSE' if sc == 404 else 'SUSPENSION'
                    with connections['myproduct'].cursor() as c:
                        c.execute("UPDATE smartstore_product SET status_type=%s WHERE id=%s", [db_st, target['id']])
                    if old_ss != db_st:
                        res['changes'].append((target, old_ss, db_st))
                    res['skipped'] += 1
                    res['skipped_items'].append({'product_code': target['product_code'], 'status': str(sc), 'store': target.get('store_name', '')})
                    time.sleep(0.3)
                    break
                else:
                    res['errors'].append({'product_code': target['product_code'], 'error': str(e)})
                    break
            except Exception as e:
                res['errors'].append({'product_code': target['product_code'], 'error': str(e)})
                break
        else:
            res['errors'].append({'product_code': target['product_code'], 'error': '429 rate limit (4회 실패)'})

    return res


def sync_status_execute(mode='all', source='api'):
    """오너클랜 sale_status → 스마트스토어 실제 동기화.
    상점별 병렬 처리 (ThreadPoolExecutor, 최대 8스레���).
    mode: 'all'|'activate'|'suspend', source: 'api'|'cron'|'manual'
    """
    import json as _json
    from concurrent.futures import ThreadPoolExecutor, as_completed

    # 중복 실행 방지 (파일 락)
    if os.path.exists(_SYNC_LOCK_PATH):
        try:
            with open(_SYNC_LOCK_PATH) as f:
                pid = int(f.read().strip())
            if os.path.exists(f'/proc/{pid}'):
                logger.info(f'[SYNC] 이미 실행 중 (pid={pid}), 스킵')
                return {'success': 0, 'fail': 0, 'skipped': 0, 'errors': [],
                        'skipped_items': [], 'detail': {'locked': True, 'pid': pid}}
        except (ValueError, IOError):
            pass
    try:
        with open(_SYNC_LOCK_PATH, 'w') as f:
            f.write(str(os.getpid()))
    except IOError:
        pass

    try:
        return _sync_status_execute_inner(mode, source)
    finally:
        try:
            os.unlink(_SYNC_LOCK_PATH)
        except OSError:
            pass


def _sync_status_execute_inner(mode, source):
    import json as _json
    from concurrent.futures import ThreadPoolExecutor, as_completed

    results = {'success': 0, 'fail': 0, 'skipped': 0, 'errors': [], 'skipped_items': [], 'detail': {}}

    suspend_targets = []
    activate_targets = []

    with connections[DB].cursor() as cur:
        if mode in ('all', 'suspend'):
            cur.execute('''
                SELECT sp.id, sp.origin_product_no, sp.store_id, op.product_code,
                       op.id as oc_id, op.sale_status, sp.status_type as old_ss_status,
                       s.commerce_api_key, s.commerce_secret_key, s.store_name
                FROM ownerclan_product op
                JOIN myproduct.smartstore_product sp ON sp.seller_management_code = op.product_code
                JOIN myproduct.smartstoreIdList s ON s.id = sp.store_id
                WHERE op.sale_status IN (2, 3) AND sp.status_type = 'SALE'
            ''')
            suspend_targets = _dictfetchall(cur)

        if mode in ('all', 'activate'):
            cur.execute('''
                SELECT sp.id, sp.origin_product_no, sp.store_id, op.product_code,
                       op.id as oc_id, op.sale_status, sp.status_type as old_ss_status,
                       s.commerce_api_key, s.commerce_secret_key, s.store_name
                FROM ownerclan_product op
                JOIN myproduct.smartstore_product sp ON sp.seller_management_code = op.product_code
                JOIN myproduct.smartstoreIdList s ON s.id = sp.store_id
                WHERE op.sale_status = 1 AND sp.status_type IN ('SUSPENSION', 'OUTOFSTOCK')
            ''')
            activate_targets = _dictfetchall(cur)

    all_targets = [
        (t, 'SUSPENSION') for t in suspend_targets
    ] + [
        (t, 'SALE') for t in activate_targets
    ]

    if not all_targets:
        return results

    # 세션 로그
    now = datetime.now()
    sync_log_id = 0
    try:
        with connections[DB].cursor() as cur:
            cur.execute(
                "INSERT INTO ownerclan_status_sync_log "
                "(started_at, activate_target, suspend_target, source) VALUES (%s,%s,%s,%s)",
                [now, len(activate_targets), len(suspend_targets), source],
            )
            sync_log_id = cur.lastrowid
    except Exception as e:
        logger.warning(f'[SYNC] 세션로그 생성 실패: {e}')

    # 상점별 그룹핑
    store_groups = {}
    for target, new_status in all_targets:
        sid = target['store_id']
        if sid not in store_groups:
            store_groups[sid] = {
                'api_key': target['commerce_api_key'],
                'secret_key': target['commerce_secret_key'],
                'items': [],
            }
        store_groups[sid]['items'].append((target, new_status))

    logger.info(f'[SYNC] {len(all_targets)}건 대상, {len(store_groups)}개 상점 병렬 처리')

    # 상점별 병렬 처리 (최대 8스레드)
    all_changes = []
    with ThreadPoolExecutor(max_workers=min(8, len(store_groups))) as pool:
        futures = {
            pool.submit(_sync_store_batch, sid, g['items'], g['api_key'], g['secret_key']): sid
            for sid, g in store_groups.items()
        }
        for future in as_completed(futures):
            res = future.result()
            results['success'] += res['success']
            results['skipped'] += res['skipped']
            results['fail'] += len(res['errors'])
            results['errors'].extend(res['errors'])
            results['skipped_items'].extend(res['skipped_items'])
            all_changes.extend(res['changes'])

    # 양쪽 로그 배치 저장
    if all_changes:
        oc_logs = []
        ss_logs = []
        changed_at = datetime.now()
        for target, old_st, new_st in all_changes:
            oc_logs.append((target.get('oc_id', 0), target['product_code'],
                            'status', 'naver_status_sync', old_st, new_st))
            ss_logs.append((target['origin_product_no'], target['product_code'],
                            'status_type', old_st, new_st, changed_at))
        try:
            with connections[DB].cursor() as cur:
                cur.executemany(
                    "INSERT INTO product_change_log "
                    "(product_id, product_code, change_group, field_name, old_value, new_value) "
                    "VALUES (%s, %s, %s, %s, %s, %s)", oc_logs)
        except Exception as e:
            logger.warning(f'[SYNC] product_change_log 저장 실패: {e}')
        try:
            with connections['myproduct'].cursor() as cur:
                cur.executemany(
                    "INSERT INTO smartstore_sync_change "
                    "(sync_log_id, origin_product_no, seller_management_code, "
                    "field_name, old_value, new_value, changed_at) "
                    "VALUES (0, %s, %s, %s, %s, %s, %s)", ss_logs)
        except Exception as e:
            logger.warning(f'[SYNC] smartstore_sync_change 저장 실패: {e}')

    # 세션 로그 업데이트
    elapsed = (datetime.now() - now).total_seconds()
    activate_ok = sum(1 for _, _, ns in all_changes if ns == 'SALE')
    suspend_ok = sum(1 for _, _, ns in all_changes if ns != 'SALE')
    if sync_log_id:
        try:
            errors_json = _json.dumps(results['errors'][:20], ensure_ascii=False) if results['errors'] else None
            with connections[DB].cursor() as cur:
                cur.execute(
                    "UPDATE ownerclan_status_sync_log SET "
                    "completed_at=%s, activate_success=%s, suspend_success=%s, "
                    "skipped=%s, failed=%s, errors=%s, elapsed_sec=%s WHERE id=%s",
                    [datetime.now(), activate_ok, suspend_ok,
                     results['skipped'], results['fail'], errors_json, round(elapsed, 1),
                     sync_log_id])
        except Exception as e:
            logger.warning(f'[SYNC] 세션로그 업데이트 실패: {e}')

    results['detail'] = {
        'suspend_targets': len(suspend_targets),
        'activate_targets': len(activate_targets),
        'sync_log_id': sync_log_id,
        'logged_changes': len(all_changes),
    }
    results['elapsed'] = round(elapsed, 1)
    return results


# ── 변경사항 로그 ──

def record_field_changes(product_id, product_code, old_row, new_row):
    """변경된 필드를 product_change_log에 기록. old_row/new_row는 dict."""
    inserts = []
    for field, ftype in CHANGE_DETECT_FIELDS.items():
        old_val = old_row.get(field)
        new_val = new_row.get(field)
        if ftype == 'int':
            old_cmp = int(old_val or 0)
            new_cmp = int(new_val or 0)
        else:
            old_cmp = str(old_val or '')
            new_cmp = str(new_val or '')
        if old_cmp != new_cmp:
            group = FIELD_TO_GROUP.get(field, 'etc')
            inserts.append((product_id, product_code, group, field,
                            str(old_val) if old_val is not None else '',
                            str(new_val) if new_val is not None else ''))
    if inserts:
        with connections[DB].cursor() as cur:
            cur.executemany(
                "INSERT INTO product_change_log "
                "(product_id, product_code, change_group, field_name, old_value, new_value) "
                "VALUES (%s, %s, %s, %s, %s, %s)",
                inserts,
            )
    return len(inserts)


def get_change_log_summary(date=None, page=1, per_page=50):
    """변경사항 통합 리포트 (필드변경 + 품절동기화)"""
    import json as _json

    with connections[DB].cursor() as cur:
        # 날짜 조건
        if date:
            date_where = "AND DATE(detected_at) = %s"
            date_params = [date]
        else:
            date_where = "AND is_applied = 0"
            date_params = []

        # 그룹별 요약
        cur.execute(f"""
            SELECT change_group, COUNT(*) as cnt,
                   COUNT(DISTINCT product_code) as product_cnt
            FROM product_change_log
            WHERE 1=1 {date_where}
            GROUP BY change_group
            ORDER BY cnt DESC
        """, date_params)
        groups = []
        total = 0
        for row in cur.fetchall():
            groups.append({
                'group': row[0],
                'label': CHANGE_GROUP_LABELS.get(row[0], row[0]),
                'count': row[1],
                'product_count': row[2],
            })
            total += row[1]

        # 페이지네이션된 상세 목록
        offset = (page - 1) * per_page
        cur.execute(f"""
            SELECT id, product_code, change_group, field_name,
                   old_value, new_value, detected_at
            FROM product_change_log
            WHERE 1=1 {date_where}
            ORDER BY detected_at DESC
            LIMIT %s OFFSET %s
        """, date_params + [per_page, offset])
        items = _dictfetchall(cur)
        for r in items:
            r['group_label'] = CHANGE_GROUP_LABELS.get(r['change_group'], r['change_group'])
            if r['detected_at']:
                r['detected_at'] = r['detected_at'].strftime('%Y-%m-%d %H:%M')

    total_pages = (total + per_page - 1) // per_page if total > 0 else 0

    # 품절동기화 이력
    soldout_sync = None
    try:
        with connections[DB].cursor() as cur:
            if date:
                cur.execute("""
                    SELECT sync_date, total_changes, status_changes,
                           transitions, db_result, elapsed
                    FROM soldout_sync_log WHERE sync_date = %s
                """, [date])
            else:
                cur.execute("""
                    SELECT sync_date, total_changes, status_changes,
                           transitions, db_result, elapsed
                    FROM soldout_sync_log ORDER BY sync_date DESC LIMIT 1
                """)
            row = cur.fetchone()
            if row:
                transitions = row[3]
                db_result = row[4]
                if isinstance(transitions, str):
                    transitions = _json.loads(transitions)
                if isinstance(db_result, str):
                    db_result = _json.loads(db_result)
                soldout_sync = {
                    'sync_date': str(row[0]),
                    'total_changes': row[1],
                    'status_changes': row[2],
                    'transitions': transitions,
                    'db_result': db_result,
                    'elapsed': row[5],
                }
    except Exception:
        pass

    return {
        'date': date or 'current',
        'field_changes': {'groups': groups, 'total': total, 'items': items},
        'soldout_sync': soldout_sync,
        'pagination': {'page': page, 'per_page': per_page, 'total_pages': total_pages},
    }


def get_change_log_dates(limit=30):
    """변경 발생 날짜 목록 (최근 30���)"""
    import json as _json
    dates = []

    with connections[DB].cursor() as cur:
        # 필드 변경 날짜별 집계
        cur.execute("""
            SELECT DATE(detected_at) as log_date, COUNT(*) as cnt
            FROM product_change_log
            GROUP BY DATE(detected_at)
            ORDER BY log_date DESC
            LIMIT %s
        """, [limit])
        field_dates = {str(row[0]): row[1] for row in cur.fetchall()}

        # 품절동기화 날짜별
        cur.execute("""
            SELECT sync_date, status_changes
            FROM soldout_sync_log
            ORDER BY sync_date DESC
            LIMIT %s
        """, [limit])
        soldout_dates = {str(row[0]): row[1] for row in cur.fetchall()}

    # 합치기
    all_dates = sorted(set(list(field_dates.keys()) + list(soldout_dates.keys())), reverse=True)[:limit]
    for d in all_dates:
        dates.append({
            'date': d,
            'field_changes': field_dates.get(d, 0),
            'soldout_changes': soldout_dates.get(d, 0),
        })

    return {'dates': dates}


def get_product_change_log(product_id):
    """특정 상품의 변경 이력"""
    with connections[DB].cursor() as cur:
        cur.execute("""
            SELECT id, product_code, change_group, field_name,
                   old_value, new_value, is_applied, detected_at, applied_at
            FROM product_change_log
            WHERE product_id = %s
            ORDER BY detected_at DESC
            LIMIT 100
        """, [product_id])
        rows = _dictfetchall(cur)
        for r in rows:
            r['group_label'] = CHANGE_GROUP_LABELS.get(r['change_group'], r['change_group'])
            if r['detected_at']:
                r['detected_at'] = r['detected_at'].strftime('%Y-%m-%d %H:%M')
            if r['applied_at']:
                r['applied_at'] = r['applied_at'].strftime('%Y-%m-%d %H:%M')
    return rows
